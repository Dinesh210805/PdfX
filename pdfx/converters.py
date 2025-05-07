"""
Converter functions for different file formats to PDF
"""
import os
from fpdf import FPDF
from PIL import Image
import docx
import pdfkit
import colorama
from colorama import Fore, Style

# Initialize colorama
colorama.init()
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PyPDF2 import PdfMerger, PdfReader, PdfWriter

def convert_txt_to_pdf(input_path, output_path=None):
    """Convert a text file to PDF"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.pdf'
    
    try:
        # Read the text file
        with open(input_path, 'r', encoding='utf-8') as file:
            content = file.read()
        
        # Create a PDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # Add text to PDF
        for line in content.split('\n'):
            pdf.cell(0, 10, txt=line, ln=True)
          # Save the PDF
        pdf.output(output_path)
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}Text file converted successfully! PDF saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error converting TXT to PDF: {str(e)}{Style.RESET_ALL}")
        return False

def convert_jpg_to_pdf(input_path, output_path=None):
    """Convert JPG/PNG image to PDF"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.pdf'
    
    try:
        # Open the image
        img = Image.open(input_path)
        
        # Convert to RGB if necessary
        if img.mode != 'RGB':
            img = img.convert('RGB')
          # Save as PDF
        img.save(output_path, "PDF", resolution=100.0)
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}Image converted successfully! PDF saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error converting image to PDF: {str(e)}{Style.RESET_ALL}")
        return False

def convert_docx_to_pdf(input_path, output_path=None):
    """Convert DOCX to PDF"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.pdf'
    
    try:
        # Load the Word document
        doc = docx.Document(input_path)
        
        # Create a PDF canvas
        c = canvas.Canvas(output_path, pagesize=letter)
        width, height = letter
        
        # Extract text and add to PDF
        for para in doc.paragraphs:
            text = para.text
            c.drawString(72, height - 72, text)
            height -= 14  # Move down for next line
            if height < 72:
                c.showPage()
                height = letter[1]
          # Save the PDF
        c.save()
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}Word document converted successfully! PDF saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error converting DOCX to PDF: {str(e)}{Style.RESET_ALL}")
        return False

def convert_html_to_pdf(input_path, output_path=None):
    """Convert HTML to PDF"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.pdf'
    
    try:        # Use pdfkit to convert HTML to PDF
        pdfkit.from_file(input_path, output_path)
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}HTML converted successfully! PDF saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error converting HTML to PDF: {str(e)}{Style.RESET_ALL}")
        return False

def convert_ppt_to_pdf(input_path, output_path=None):
    """Convert PPT/PPTX to PDF"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.pdf'
    
    try:
        import comtypes.client
        
        # Create PowerPoint application
        powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
        powerpoint.Visible = 1
        
        # Open the presentation
        presentation = powerpoint.Presentations.Open(os.path.abspath(input_path))
        
        # Save as PDF
        presentation.SaveAs(os.path.abspath(output_path), 32)  # 32 = PDF format
          # Close
        presentation.Close()
        powerpoint.Quit()
        
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}PowerPoint presentation converted successfully! PDF saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error converting PPT to PDF: {str(e)}{Style.RESET_ALL}")
        return False

def merge_pdfs(input_paths, output_path):
    """Merge multiple PDFs into one"""
    try:
        merger = PdfMerger()
        
        # Add each PDF to the merger
        for pdf in input_paths:
            merger.append(pdf)
          # Write to the output file
        merger.write(output_path)
        merger.close()
        
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}PDFs merged successfully! Output saved to: {output_path}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error merging PDFs: {str(e)}{Style.RESET_ALL}")
        return False

def split_pdf(input_path, output_directory=None):
    """Split a PDF into individual pages"""
    if not output_directory:
        output_directory = os.path.dirname(input_path)
    
    try:
        # Read the input PDF
        pdf = PdfReader(input_path)
        
        # Get the base filename
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        
        # Split the PDF
        for i, page in enumerate(pdf.pages):
            output_path = os.path.join(output_directory, f"{base_name}_page_{i+1}.pdf")
            writer = PdfWriter()
            writer.add_page(page)
            with open(output_path, "wb") as output_file:
                writer.write(output_file)
        
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}PDF split successfully! {len(pdf.pages)} pages extracted to: {output_directory}{Style.RESET_ALL}")
        return True
    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error splitting PDF: {str(e)}{Style.RESET_ALL}")
        return False

def extract_text_from_pdf(input_path, output_path=None):
    """Extract text from a PDF file"""
    if not output_path:
        output_path = os.path.splitext(input_path)[0] + '.txt'
    
    try:
        # Read the PDF
        reader = PdfReader(input_path)
        
        # Extract text from each page
        with open(output_path, 'w', encoding='utf-8') as output_file:
            for page in reader.pages:
                text = page.extract_text()
                output_file.write(text + "\n\n")
        print(f"\n{Fore.WHITE}[{Fore.GREEN}+{Fore.WHITE}] {Fore.GREEN}Text extracted successfully! Saved to: {output_path}{Style.RESET_ALL}")
        return True

    except Exception as e:
        print(f"\n{Fore.WHITE}[{Fore.RED}!{Fore.WHITE}] {Fore.RED}Error extracting text from PDF: {str(e)}{Style.RESET_ALL}")
        return False
